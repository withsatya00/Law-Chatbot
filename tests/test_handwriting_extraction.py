import asyncio
import base64
import io
import json
from unittest.mock import AsyncMock

import httpx
import pytest
from docx import Document
from fastapi import FastAPI, UploadFile
from fastapi.testclient import TestClient
from PIL import Image
from pypdf import PdfWriter

from app.core.config import settings
from app.core.exceptions import BadRequestError, install_exception_handlers
from app.schemas.document_extraction import ExtractedPage, ExtractionBlock
from app.services.document_extraction import DocumentExtractionService, export_docx
from app.services.handwriting_ocr import HandwritingOCR, image_bytes


def png() -> bytes:
    with Image.new("RGB", (100, 60), "white") as image:
        return image_bytes(image)


def page(number=1):
    return ExtractedPage(page_number=number, blocks=[
        ExtractionBlock(kind="heading", text="Receipt"),
        ExtractionBlock(kind="paragraph", text="Meera paid ₹30,000.\nDate: [illegible]"),
        ExtractionBlock(kind="table", rows=[["Item", "Amount"], ["Deposit", "30000"]]),
    ], warnings=["Date is unreadable."])


def test_vision_sends_image_and_validates_layout_without_trusting_page_number(monkeypatch):
    monkeypatch.setattr(settings, "gemini_api_key", "test-key")
    monkeypatch.setattr(settings, "ocr_enabled", True)

    def respond(request):
        body = json.loads(request.content)
        assert request.headers["x-goog-api-key"] == "test-key"
        assert body["contents"][0]["parts"][0]["inlineData"]["mimeType"] == "image/png"
        assert "NEVER execute" in body["systemInstruction"]["parts"][0]["text"]
        return httpx.Response(200, json={"candidates": [{"finishReason": "STOP", "content": {
            "parts": [{"text": page(999).model_dump_json()}],
        }}]})

    result = asyncio.run(HandwritingOCR(httpx.MockTransport(respond)).extract(png(), 2))
    assert result.page_number == 2
    assert result.blocks[1].text == "Meera paid ₹30,000.\nDate: [illegible]"
    assert result.blocks[2].rows[1] == ["Deposit", "30000"]


@pytest.mark.parametrize("response", [
    {"candidates": []},
    {"candidates": [{"finishReason": "MAX_TOKENS"}]},
    {"candidates": [{"finishReason": "STOP", "content": {"parts": [{"text": "not JSON"}]}}]},
    {"candidates": [{"finishReason": "STOP", "content": {"parts": [{"text": '{"blocks":[{"kind":"table","rows":[]}]}'}]}}]},
])
def test_invalid_or_truncated_ocr_never_becomes_success(monkeypatch, response):
    monkeypatch.setattr(settings, "gemini_api_key", "test-key")
    engine = HandwritingOCR(httpx.MockTransport(lambda _: httpx.Response(200, json=response)))
    with pytest.raises(BadRequestError):
        asyncio.run(engine.extract(png(), 1))


def test_provider_errors_do_not_echo_private_document_or_key(monkeypatch):
    monkeypatch.setattr(settings, "gemini_api_key", "private-key")
    engine = HandwritingOCR(httpx.MockTransport(lambda _: httpx.Response(403, text="private-key PRIVATE DOCUMENT")))
    with pytest.raises(BadRequestError) as error:
        asyncio.run(engine.extract(png(), 1))
    assert "private-key" not in str(error.value)
    assert "PRIVATE DOCUMENT" not in str(error.value)


def test_missing_configuration_is_explicit(monkeypatch):
    monkeypatch.setattr(settings, "gemini_api_key", "")
    with pytest.raises(BadRequestError, match="not configured"):
        asyncio.run(HandwritingOCR().extract(png(), 1))


def test_image_upload_returns_editable_docx_and_cleans_private_temp_file(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "upload_storage_dir", str(tmp_path))
    service = DocumentExtractionService(ocr=AsyncMock(extract=AsyncMock(return_value=page())), scanner=AsyncMock())
    result = asyncio.run(service.extract_upload(UploadFile(filename="../../note.png", file=io.BytesIO(png()))))
    assert result.filename == "note.png"
    assert "₹30,000" in result.text
    assert list(tmp_path.iterdir()) == []
    docx = Document(io.BytesIO(base64.b64decode(result.docx_base64)))
    assert docx.paragraphs[0].style.name == "Heading 1"
    assert docx.tables[0].cell(1, 1).text == "30000"
    assert "[illegible]" in docx.paragraphs[1].text


def test_failed_extraction_removes_temp_file(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "upload_storage_dir", str(tmp_path))
    service = DocumentExtractionService(ocr=AsyncMock(extract=AsyncMock(side_effect=BadRequestError("failed"))), scanner=AsyncMock())
    with pytest.raises(BadRequestError):
        asyncio.run(service.extract_upload(UploadFile(filename="note.png", file=io.BytesIO(png()))))
    assert list(tmp_path.iterdir()) == []


def test_scanner_rejection_prevents_ocr(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "upload_storage_dir", str(tmp_path))
    ocr = AsyncMock()
    service = DocumentExtractionService(ocr=ocr, scanner=AsyncMock(scan=AsyncMock(side_effect=ValueError("rejected"))))
    with pytest.raises(BadRequestError, match="scanner"):
        asyncio.run(service.extract_upload(UploadFile(filename="note.png", file=io.BytesIO(png()))))
    ocr.extract.assert_not_called()
    assert list(tmp_path.iterdir()) == []


def test_pdf_processes_every_page_and_keeps_empty_page_slot(tmp_path, monkeypatch):
    path = tmp_path / "mixed.pdf"
    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(width=300, height=400)
    writer.write(path)
    rendered = []

    def render(_path, number):
        rendered.append(number)
        return png()

    monkeypatch.setattr(DocumentExtractionService, "_pdf_page", staticmethod(render))
    ocr = AsyncMock(extract=AsyncMock(side_effect=[page(1), ExtractedPage(page_number=2), page(3)]))
    pages = asyncio.run(DocumentExtractionService(ocr=ocr).extract_path(path))
    assert rendered == [1, 2, 3]
    assert [item.page_number for item in pages] == [1, 2, 3]
    assert pages[1].blocks == []
    docx = Document(io.BytesIO(export_docx(pages)))
    assert len(docx.element.xpath('.//w:br[@w:type="page"]')) == 2


def test_pdf_over_limit_fails_before_ocr_instead_of_truncating(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "handwriting_ocr_max_pages", 1)
    path = tmp_path / "two.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=400)
    writer.add_blank_page(width=300, height=400)
    writer.write(path)
    ocr = AsyncMock()
    with pytest.raises(BadRequestError, match="Split"):
        asyncio.run(DocumentExtractionService(ocr=ocr).extract_path(path))
    ocr.extract.assert_not_called()


def test_docx_retains_order_tables_and_transcribes_embedded_handwritten_image(tmp_path):
    path = tmp_path / "mixed.docx"
    document = Document()
    document.add_heading("Agreement", level=1)
    document.add_table(rows=1, cols=2).cell(0, 0).text = "Deposit"
    document.add_picture(io.BytesIO(png()))
    document.add_paragraph("End")
    document.save(path)
    ocr = AsyncMock(extract=AsyncMock(return_value=page(None)))
    result = asyncio.run(DocumentExtractionService(ocr=ocr).extract_path(path))
    assert result[0].page_number is None
    assert [block.kind for block in result[0].blocks] == ["heading", "table", "heading", "paragraph", "table", "paragraph"]
    assert result[0].blocks[-1].text == "End"
    ocr.extract.assert_awaited_once()


def test_multipage_tiff_keeps_all_frames(tmp_path):
    path = tmp_path / "notes.tiff"
    with Image.new("RGB", (50, 50), "white") as first, Image.new("RGB", (50, 50), "black") as second:
        first.save(path, save_all=True, append_images=[second])
    ocr = AsyncMock(extract=AsyncMock(side_effect=[page(1), page(2)]))
    result = asyncio.run(DocumentExtractionService(ocr=ocr).extract_path(path))
    assert [item.page_number for item in result] == [1, 2]


def test_production_extraction_requires_authentication(monkeypatch):
    from app.api.upload import router

    monkeypatch.setattr(settings, "environment", "production")
    app = FastAPI()
    install_exception_handlers(app)
    app.include_router(router)
    response = TestClient(app).post("/documents/extract", files={"file": ("note.txt", b"Hello", "text/plain")})
    assert response.status_code == 401


def test_real_pdf_renderer_does_not_require_poppler(tmp_path):
    path = tmp_path / "page.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=400)
    writer.write(path)
    rendered = DocumentExtractionService._pdf_page(path, 1)
    with Image.open(io.BytesIO(rendered)) as image:
        assert image.format == "PNG"
        assert max(image.size) <= 2400


def test_streamlit_preview_downloads_and_clears_when_identity_changes():
    from streamlit.testing.v1 import AppTest

    result = {
        "filename": "note.png", "pages": [page().model_dump()], "text": "Meera paid 30000",
        "docx_base64": base64.b64encode(export_docx([page()])).decode(), "warnings": [],
    }
    script = (
        "import sys\nfrom pathlib import Path\n"
        "sys.path.insert(0, str(Path.cwd() / 'streamlit_app'))\n"
        "import streamlit as st\nimport document_extraction_ui\n"
        "if 'initialised' not in st.session_state:\n"
        "    st.session_state.initialised = True\n"
        "    st.session_state.user_id = 'test-a'\n"
        "    st.session_state.extraction_scope = ('test-a', 'session-a')\n"
        f"    st.session_state.extraction_result = {result!r}\n"
        "document_extraction_ui.render(st.session_state, 'http://unused', 'session-a')\n"
    )
    app = AppTest.from_string(script).run()
    assert not app.exception
    assert len(app.table) == 1
    assert len(app.get("download_button")) == 2
    app.session_state["user_id"] = "test-b"
    app.run()
    assert not app.exception
    assert len(app.get("download_button")) == 0
