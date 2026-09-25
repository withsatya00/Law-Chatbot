"""Advocate-register redesign (see app/drafting/advocate_register.py):
gating, grammatical degradation when optional fields are blank, and --
critically -- that the new hedged "typical supporting documents" sentence
never trips `prohibited_clauses.py`'s `evidence_not_supplied` rule, which
exists to stop a draft from claiming documents are attached when they are
not.
"""

import asyncio
from unittest.mock import AsyncMock

from app.drafting import advocate_register, prohibited_clauses
from app.drafting.engine import LegalDraftEngine
from app.drafting.export import _LAYOUT, DocxDraftExporter, ExportOptions, PdfDraftExporter
from app.drafting.templates import get_template, list_templates
from app.drafting.templates.base import DraftTemplateDefinition
from app.llm.base import LLMResponse
from app.schemas.drafting import DraftGenerateRequest, DraftGenerateResponse


def _generate(draft_id: str, fields: dict[str, str]) -> DraftGenerateResponse:
    engine = LegalDraftEngine()
    engine.llm.chat = AsyncMock(return_value=LLMResponse(content="", model="test", provider="test"))
    return asyncio.run(engine.preview(DraftGenerateRequest(draft_id=draft_id, fields=fields)))


def _template(**overrides) -> DraftTemplateDefinition:
    base = {
        "draft_id": "test_template", "name": "Test", "hindi_name": "परीक्षण", "category": "Complaint",
        "description": "A test template.", "authority_label": "The Authority,",
        "applicable_acts_hint": [], "applicable_sections_hint": [], "required_fields": [],
    }
    base.update(overrides)
    return DraftTemplateDefinition(**base)


# ---------------------------------------------------------------------------
# Gating
# ---------------------------------------------------------------------------


def test_gating_covers_english_hindi_hinglish_only() -> None:
    template = _template()
    assert advocate_register.supports_advocate_register("english", template)
    assert advocate_register.supports_advocate_register("hindi", template)
    assert advocate_register.supports_advocate_register("hinglish", template)
    assert not advocate_register.supports_advocate_register("tamil", template)
    assert not advocate_register.supports_advocate_register("bengali", template)


def test_gating_covers_complaint_and_application_only() -> None:
    assert advocate_register.supports_advocate_register("english", _template(category="Complaint"))
    assert advocate_register.supports_advocate_register("english", _template(category="Application"))
    assert not advocate_register.supports_advocate_register("english", _template(category="Notice"))
    assert not advocate_register.supports_advocate_register("english", _template(category="Affidavit"))
    assert not advocate_register.supports_advocate_register("english", _template(category="Contract"))


def test_ordinary_correspondence_application_templates_opt_out() -> None:
    # Confirmed live in templates/*.yaml: resignation_letter, job_application,
    # etc. set advocate_register: false because the sworn register reads
    # wrong for ordinary employer/school correspondence.
    template = _template(category="Application", advocate_register=False)
    assert not advocate_register.supports_advocate_register("english", template)


def test_real_ordinary_correspondence_templates_have_the_flag_set() -> None:
    for draft_id in (
        "resignation_letter", "job_application", "education_leave_application",
        "employment_leave_application", "employment_noc_request", "experience_certificate_request",
        "scholarship_application", "tc_application",
    ):
        template = get_template(draft_id)
        assert template is not None, draft_id
        assert template.advocate_register is False, draft_id


def test_real_formal_petition_templates_keep_the_default() -> None:
    for draft_id in ("rti_application", "income_certificate_application", "police_complaint", "mobile_theft_complaint"):
        template = get_template(draft_id)
        assert template is not None, draft_id
        assert template.advocate_register is True, draft_id


# ---------------------------------------------------------------------------
# sworn_opening_block -- grammatical degradation
# ---------------------------------------------------------------------------


def test_sworn_opening_full_fields_english() -> None:
    fields = {
        "applicant_name": "Sunita Devi", "applicant_father_name": "Ramesh Sharma",
        "applicant_age": "35", "applicant_address": "H.No. 12, Lucknow", "applicant_mobile": "9876543210",
    }
    text = advocate_register.sworn_opening_block(fields, "english")
    assert text.startswith("I, Sunita Devi, S/o, D/o, W/o Ramesh Sharma, aged about 35 years, "
                            "residing at H.No. 12, Lucknow, and holding Mobile Number 9876543210, "
                            "do hereby state as under:-")


def test_sworn_opening_degrades_gracefully_when_new_fields_are_blank() -> None:
    # The realistic case for every existing draft today: age/father_name are
    # brand-new fields nobody has filled in yet.
    fields = {"applicant_name": "Neeraj Singh", "applicant_address": "Ghaziabad", "applicant_mobile": "9012345678"}
    text = advocate_register.sworn_opening_block(fields, "english")
    assert text == "I, Neeraj Singh, residing at Ghaziabad, and holding Mobile Number 9012345678, do hereby state as under:-"
    # No placeholder blanks, no leftover "aged about  years", no double commas.
    assert "aged" not in text
    assert ",," not in text
    assert "  " not in text


def test_sworn_opening_degrades_gracefully_when_address_is_also_blank() -> None:
    # bonafide_certificate_application never asks for an address.
    fields = {"applicant_name": "Aditi Rao", "applicant_mobile": "9012345678"}
    text = advocate_register.sworn_opening_block(fields, "english")
    assert text == "I, Aditi Rao, holding Mobile Number 9012345678, do hereby state as under:-"
    assert ",," not in text


def test_sworn_opening_with_only_name() -> None:
    text = advocate_register.sworn_opening_block({"applicant_name": "Aditi Rao"}, "english")
    assert text == "I, Aditi Rao, do hereby state as under:-"


def test_sworn_opening_hindi_full_fields() -> None:
    fields = {
        "applicant_name": "नीरज सिंह", "applicant_father_name": "सुरेश सिंह",
        "applicant_age": "29", "applicant_address": "गाज़ियाबाद", "applicant_mobile": "9012345678",
    }
    text = advocate_register.sworn_opening_block(fields, "hindi")
    assert text.startswith("मैं, नीरज सिंह,")
    assert text.endswith("निम्नानुसार कथन करता/करती हूँ:-")
    assert ",," not in text


def test_sworn_opening_no_name_returns_not_provided_not_a_crash() -> None:
    text = advocate_register.sworn_opening_block({}, "english")
    assert text == "Not provided."


# ---------------------------------------------------------------------------
# classical_prayer_intro / append_citation -- never invents a citation
# ---------------------------------------------------------------------------


def test_classical_prayer_intro_none_when_template_has_no_citation() -> None:
    assert advocate_register.classical_prayer_intro("english", [], []) is None


def test_classical_prayer_intro_present_when_hints_exist() -> None:
    text = advocate_register.classical_prayer_intro("english", ["Section 303 BNS"], ["Bharatiya Nyaya Sanhita"])
    assert text is not None
    assert "most respectfully prayed" in text
    assert "Section 303 BNS" in text


def test_append_citation_is_a_no_op_with_no_hints() -> None:
    assert advocate_register.append_citation("Complaint regarding theft", "english", [], []) == "Complaint regarding theft"


def test_append_citation_appends_under_clause() -> None:
    subject = advocate_register.append_citation("Complaint regarding theft", "english", ["Section 303 BNS"], [])
    assert subject == "Complaint regarding theft under Section 303 BNS"


def test_append_citation_is_a_no_op_when_subject_already_names_the_section() -> None:
    # cheque_bounce_notice.yaml's own subject_template already names Section
    # 138 -- confirmed live, appending unconditionally produced "...under
    # Section 138... under Section 138, Section 142" (the same citation
    # twice in one line).
    subject = "Statutory notice under Section 138, Negotiable Instruments Act, 1881 -- dishonour of cheque no. 123456"
    result = advocate_register.append_citation(subject, "english", ["Section 138", "Section 142"], [])
    assert result == subject


# ---------------------------------------------------------------------------
# supports_citation -- broader than supports_advocate_register (Notice too)
# ---------------------------------------------------------------------------


def test_supports_citation_covers_notice_too() -> None:
    # Notice does NOT get the sworn register, but DOES get Subject-line
    # citation -- confirmed against real cheque-bounce/money-recovery notice
    # conventions (see advocate_register.py's module docstring).
    template = _template(category="Notice")
    assert advocate_register.supports_citation("english", template)
    assert not advocate_register.supports_advocate_register("english", template)


def test_supports_citation_excludes_affidavit_and_contract() -> None:
    assert not advocate_register.supports_citation("english", _template(category="Affidavit"))
    assert not advocate_register.supports_citation("english", _template(category="Contract"))


# ---------------------------------------------------------------------------
# enclosures_block -- the fabrication-risk case
# ---------------------------------------------------------------------------


def test_enclosures_prefers_user_supplied_documents() -> None:
    template = _template(typical_supporting_documents=["Should never appear"])
    text = advocate_register.enclosures_block({"available_documents": "RC copy, insurance papers"}, template, "english")
    assert text == "RC copy, insurance papers"


def test_enclosures_falls_back_to_hedged_suggestion_when_nothing_supplied() -> None:
    template = _template(typical_supporting_documents=["Photocopy of Aadhaar Card"])
    text = advocate_register.enclosures_block({}, template, "english")
    assert text is not None
    assert "Photocopy of Aadhaar Card" in text
    assert "ordinarily required" in text
    assert "should be enclosed where available" in text


def test_enclosures_none_when_nothing_supplied_and_no_typical_documents() -> None:
    template = _template()
    assert advocate_register.enclosures_block({}, template, "english") is None


def test_enclosures_hedge_never_trips_evidence_not_supplied_rule() -> None:
    # The single most important test in this file: prohibited_clauses.py's
    # evidence_not_supplied rule exists specifically to catch a draft
    # claiming documents are attached when the user listed none. The
    # typical-documents hedge sentence must never trigger it.
    template = _template(typical_supporting_documents=["Photocopy of Aadhaar Card", "Photocopy of RC"])
    text = advocate_register.enclosures_block({}, template, "english")
    assert text is not None
    findings = prohibited_clauses.scan({"Enclosures": text}, fields={})
    evidence_findings = [f for f in findings if f.get("category") == "unsupported_statement:evidence_not_supplied"]
    assert evidence_findings == [], f"hedge sentence tripped evidence_not_supplied: {findings!r}"


# ---------------------------------------------------------------------------
# End-to-end: a real Notice template gets the citation, never the sworn
# register or Enclosures -- both of those stay Complaint/Application-only.
# ---------------------------------------------------------------------------


def test_tenancy_termination_notice_subject_cites_section_106() -> None:
    fields = {
        "applicant_name": "Landlord X", "applicant_address": "Y", "applicant_mobile": "9876543210",
        "respondent_name": "Tenant Z", "rented_premises_address": "Flat 5, ABC Apartments, Delhi",
        "facts": "The tenancy is a month-to-month arrangement and I wish to terminate it.", "place": "Delhi",
    }
    response = _generate("tenancy_termination_notice", fields)
    assert "Section 106, Transfer of Property Act, 1882" in response.sections["Subject"]


def test_notice_sender_details_stays_plain_not_sworn() -> None:
    fields = {
        "applicant_name": "Landlord X", "applicant_address": "Y", "applicant_mobile": "9876543210",
        "respondent_name": "Tenant Z", "rented_premises_address": "Flat 5, ABC Apartments, Delhi",
        "facts": "The tenancy is a month-to-month arrangement and I wish to terminate it.", "place": "Delhi",
    }
    response = _generate("tenancy_termination_notice", fields)
    sender_details = response.sections["Sender Details"]
    assert "do hereby state as under" not in sender_details
    assert sender_details == "Landlord X\nY\nMobile: 9876543210"


def test_notice_never_gets_an_enclosures_heading() -> None:
    # Notice keeps "Annexures" (or nothing) -- "Enclosures" is
    # Complaint/Application-only.
    fields = {
        "applicant_name": "Landlord X", "applicant_address": "Y", "applicant_mobile": "9876543210",
        "respondent_name": "Tenant Z", "rented_premises_address": "Flat 5, ABC Apartments, Delhi",
        "facts": "The tenancy is a month-to-month arrangement and I wish to terminate it.", "place": "Delhi",
        "available_documents": "Copy of the rent agreement",
    }
    response = _generate("tenancy_termination_notice", fields)
    assert "Enclosures" not in response.sections
    assert response.sections.get("Annexures") == "Copy of the rent agreement"


def test_cheque_bounce_notice_subject_does_not_duplicate_its_own_citation() -> None:
    fields = {
        "applicant_name": "Rahul Kumar Sharma", "applicant_address": "24, Sector 62, Noida",
        "applicant_mobile": "9876543210", "respondent_name": "Mangal Prasad Verma",
        "respondent_address": "15, Civil Lines, Noida", "cheque_number": "123456", "cheque_amount": "50000",
        "cheque_date": "2026-08-01", "bank_name": "State Bank of India", "dishonour_reason": "Insufficient funds",
        "dishonour_date": "2026-08-20",
        "facts": "The cheque was issued towards repayment of a loan and was dishonoured for insufficient funds.",
        "expected_relief": "Pay the cheque amount within 15 days of receipt of this notice.", "place": "Noida",
    }
    response = _generate("cheque_bounce_notice", fields)
    subject = response.sections["Subject"]
    assert subject.count("Section 138") == 1


# ---------------------------------------------------------------------------
# Flowing-letter format: the 9 SHO-representation Complaint templates render
# as one continuous letter, no internal headings -- confirmed against a real
# advocate-drafted FIR registration application (see
# advocate_register.FLOWING_LETTER_DRAFTS's own docstring). Consumer-forum
# complaints (consumer_complaint/ecommerce_complaint/service_complaint) are
# explicitly EXCLUDED -- research confirmed a real District Consumer Disputes
# Redressal Forum complaint is even MORE formally headed than this app's
# normal output, since it's a pleading before a tribunal, not a letter.
# ---------------------------------------------------------------------------


def test_flowing_letter_drafts_are_exactly_the_sho_style_complaints() -> None:
    expected = {
        "bank_fraud_complaint", "cyber_crime_complaint", "missing_person_report",
        "mobile_theft_complaint", "online_fraud_complaint", "police_complaint",
        "sp_complaint", "threat_complaint", "vehicle_theft_complaint",
    }
    assert advocate_register.FLOWING_LETTER_DRAFTS == expected
    for draft_id in expected:
        template = get_template(draft_id)
        assert template is not None, draft_id
        assert template.category == "Complaint", draft_id


def test_consumer_forum_complaints_are_not_flowing_letters() -> None:
    for draft_id in ("consumer_complaint", "ecommerce_complaint", "service_complaint"):
        assert draft_id not in advocate_register.FLOWING_LETTER_DRAFTS


def test_every_flowing_letter_draft_is_a_real_complaint_template() -> None:
    # Guards against a typo in the set silently referring to a draft_id that
    # doesn't exist.
    all_ids = {t.draft_id for t in list_templates()}
    assert advocate_register.FLOWING_LETTER_DRAFTS <= all_ids


_FLOWING_SECTIONS = {
    "Recipient": "The Station House Officer,\nKamla Nagar, Agra",
    "Subject": "Complaint regarding theft under Section 303 BNS",
    "Complainant Details": "I, Amit Kumar Sharma, residing at Agra, do hereby state as under:-",
    "Introduction": "Sir/Madam,\n\nThe applicant respectfully submits this document.",
    "Facts of the Case": "1. That the incident occurred on 5 September 2026.",
    "Legal Position": "The facts prima facie attract the provisions of the applicable law.",
    "Consequences": "No separate loss is asserted.",
    "Request": "It is therefore prayed that the FIR be registered.",
    "Verification": "Verified that the contents are true and correct.",
    "Enclosures": "1. Photocopy of Aadhaar Card",
    "Signature Block": "Place: Agra\nDate: 10 September 2026\n\nYours faithfully,\n\n(Amit Kumar Sharma)",
}


def test_flowing_pdf_html_has_no_body_section_headings() -> None:
    html = PdfDraftExporter()._build_html(
        "Police Complaint", _FLOWING_SECTIONS, _LAYOUT, "english", ExportOptions(flowing_letter=True)
    )
    for heading in ("Facts of the Case", "Legal Position", "Consequences", "Request", "Verification"):
        assert f"<h2 class='heading'>{heading}</h2>" not in html, heading
    # Recipient/Subject also lose their heading -- Recipient becomes a plain
    # "To," paragraph, Subject an inline bold-label paragraph.
    assert "<h2 class='heading'>Recipient</h2>" not in html
    assert "<h2 class='heading'>Subject</h2>" not in html
    assert "To,<br/>The Station House Officer" in html
    assert "<strong>Subject:</strong>" in html


def test_hindi_flowing_pdf_localizes_document_chrome() -> None:
    html = PdfDraftExporter()._build_html(
        "सामान्य कानूनी नोटिस", _FLOWING_SECTIONS, _LAYOUT, "hindi",
        ExportOptions(flowing_letter=True, include_header_footer=True),
    )
    assert "प्रति,<br/>The Station House Officer" in html
    assert "<strong>विषय:</strong>" in html
    assert "'पृष्ठ ' counter(page) ' का ' counter(pages)" in html
    assert "To,<br/>" not in html
    assert "<strong>Subject:</strong>" not in html


def test_flowing_pdf_html_keeps_enclosures_heading_but_not_signature() -> None:
    html = PdfDraftExporter()._build_html(
        "Police Complaint", _FLOWING_SECTIONS, _LAYOUT, "english", ExportOptions(flowing_letter=True)
    )
    assert "<h2 class='heading'>Enclosures</h2>" in html
    assert "<h2 class='heading'>Signature Block</h2>" not in html
    assert "signature-section" in html  # still gets the right-aligned closing treatment


def test_flowing_pdf_html_drops_no_content() -> None:
    html = PdfDraftExporter()._build_html(
        "Police Complaint", _FLOWING_SECTIONS, _LAYOUT, "english", ExportOptions(flowing_letter=True)
    )
    for text in _FLOWING_SECTIONS.values():
        first_line = text.splitlines()[0]
        assert first_line.split(":")[-1].strip()[:20] in html or first_line[:20] in html, first_line


def test_non_flowing_pdf_html_is_unaffected() -> None:
    # Regression guard: the default (flowing_letter=False) must still
    # produce the normal headed layout.
    html = PdfDraftExporter()._build_html("Police Complaint", _FLOWING_SECTIONS, _LAYOUT, "english", ExportOptions())
    assert "<h2 class='heading'>Facts of the Case</h2>" in html
    assert "<h2 class='heading'>Request</h2>" in html


def test_flowing_docx_has_no_body_section_headings(tmp_path) -> None:
    from docx import Document

    output_path = tmp_path / "flowing.docx"
    DocxDraftExporter().export(
        "Police Complaint", _FLOWING_SECTIONS, output_path, "english", ExportOptions(flowing_letter=True)
    )
    document = Document(str(output_path))
    heading_texts = [p.text for p in document.paragraphs if p.style.name == "Heading 2"]
    # Only Enclosures keeps a heading in flowing mode.
    assert heading_texts == ["Enclosures"]
    # The document title IS kept (the big "POLICE COMPLAINT" heading) --
    # only the internal body-section headings are removed. The small
    # per-page running header duplicating this same title is what's
    # suppressed instead, see `test_flowing_pdf_html_suppresses_running_header`.
    assert any(p.style.name == "Heading 1" and p.text == "POLICE COMPLAINT" for p in document.paragraphs)
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert full_text.startswith("POLICE COMPLAINT\nTo,\nThe Station House Officer")
    assert "Subject: Complaint regarding theft" in full_text


def test_flowing_pdf_html_keeps_the_title() -> None:
    html = PdfDraftExporter()._build_html(
        "Police Complaint", _FLOWING_SECTIONS, _LAYOUT, "english", ExportOptions(flowing_letter=True)
    )
    assert "<h1 class='title'>POLICE COMPLAINT</h1>" in html


def test_flowing_pdf_html_suppresses_running_header_but_keeps_page_footer() -> None:
    # The small per-page "@top-center" header duplicated the title on every
    # page -- suppressed; the bottom page-number footer is untouched.
    options = ExportOptions(flowing_letter=True, include_header_footer=True)
    html = PdfDraftExporter()._build_html("Police Complaint", _FLOWING_SECTIONS, _LAYOUT, "english", options)
    assert "@top-center { content: ''" in html
    assert "@bottom-center { content: 'Page '" in html


def test_non_flowing_pdf_html_also_suppresses_the_running_header() -> None:
    # User-reported layout feedback: the small duplicate "POLICE COMPLAINT"
    # heading at the top of every page (above the document's own centered
    # title heading) is gone for every template now, not just the 9
    # `FLOWING_LETTER_DRAFTS` ones -- only `h1.title` remains as the one
    # heading a reader sees. The page-number footer is unaffected.
    options = ExportOptions(include_header_footer=True)
    html = PdfDraftExporter()._build_html("Police Complaint", _FLOWING_SECTIONS, _LAYOUT, "english", options)
    assert "@top-center { content: ''" in html
    assert "@bottom-center { content: 'Page '" in html


def test_non_flowing_docx_still_has_a_title(tmp_path) -> None:
    from docx import Document

    output_path = tmp_path / "headed.docx"
    DocxDraftExporter().export("Police Complaint", _FLOWING_SECTIONS, output_path, "english", ExportOptions())
    document = Document(str(output_path))
    assert any(p.style.name == "Heading 1" and p.text == "POLICE COMPLAINT" for p in document.paragraphs)
